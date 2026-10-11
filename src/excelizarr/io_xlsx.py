"""The ``.xlsx`` side: read a workbook into a :class:`~.model.WorkbookModel`, and back.

Two readers:

- ``engine="openpyxl"`` (default) is the full-fidelity one. It opens the file twice:
  once with ``data_only=False`` for formula text, number formats, merges, column widths
  and defined names, and once with ``data_only=True`` for the *cached* results of those
  formulas, which is what pandas would have shown you.
- ``engine="calamine"`` (``python-calamine``, the optional ``fast`` extra) reads cell
  values only, much faster. No formulas, no formats, no merges.

There is one writer, through openpyxl.
"""

from __future__ import annotations

import os
from typing import Any, Literal

import numpy as np

from .kinds import CellKind, Slot
from .model import NumberFormats, SheetGrid, WorkbookModel

__all__ = ["dump_xlsx", "load_xlsx"]

Engine = Literal["openpyxl", "calamine"]
PathLike = str | os.PathLike[str]


def load_xlsx(path: PathLike, engine: Engine = "openpyxl") -> WorkbookModel:
    """Read an ``.xlsx`` file into a :class:`~.model.WorkbookModel`.

    Parameters
    ----------
    path : str or os.PathLike
        The workbook to read.
    engine : {"openpyxl", "calamine"}
        ``openpyxl`` keeps formulas, number formats and layout; ``calamine`` reads
        values only, much faster.

    Returns
    -------
    WorkbookModel
        Every sheet, in tab order.

    Raises
    ------
    ValueError
        If ``engine`` is not one of the two above.
    """
    if engine == "openpyxl":
        return _load_openpyxl(path)
    if engine == "calamine":
        return _load_calamine(path)
    raise ValueError(f"unknown engine {engine!r}; expected 'openpyxl' or 'calamine'")


def _load_openpyxl(path: PathLike) -> WorkbookModel:
    import openpyxl

    formulas_wb = openpyxl.load_workbook(path, data_only=False)
    values_wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        numfmts = NumberFormats()
        sheets = []
        for ws in formulas_wb.worksheets:
            sheets.append(_sheet_openpyxl(ws, values_wb[ws.title], numfmts))
        defined = {
            name: str(dn.attr_text)
            for name, dn in formulas_wb.defined_names.items()
            if dn.attr_text is not None
        }
    finally:
        values_wb.close()
    return WorkbookModel(sheets, numfmts, defined)


def _sheet_openpyxl(ws: Any, values_ws: Any, numfmts: NumberFormats) -> SheetGrid:
    from openpyxl.worksheet.formula import ArrayFormula

    # ``max_row``/``max_column`` are 1 on an empty sheet, so ask whether anything is
    # actually there before trusting them.
    has_cells = bool(ws._cells)
    shape = (ws.max_row, ws.max_column) if has_cells else (0, 0)
    grid = SheetGrid.empty(ws.title, shape)
    kind = grid.arrays[Slot.KIND]
    formula = grid.arrays[Slot.FORMULA]
    numfmt = grid.arrays[Slot.NUMFMT]

    if has_cells:
        values_ws.reset_dimensions()
        for i, row in enumerate(
            values_ws.iter_rows(max_row=shape[0], max_col=shape[1])
        ):
            for cell in row:
                if cell.value is None:
                    continue
                j = cell.column - 1
                value = cell.value
                if cell.data_type == "e":
                    grid.arrays[Slot.STRING][i, j] = str(value)
                    kind[i, j] = CellKind.ERROR
                else:
                    grid.set_cell(i, j, value)

        for (r, c), cell in ws._cells.items():
            i, j = r - 1, c - 1
            if cell.data_type == "f":
                f = cell.value
                formula[i, j] = f.text if isinstance(f, ArrayFormula) else str(f)
            if cell.has_style and cell.number_format != "General":
                numfmt[i, j] = numfmts.index(cell.number_format)

    grid.merged = [str(r) for r in ws.merged_cells.ranges]
    grid.col_widths = {
        letter: float(dim.width)
        for letter, dim in ws.column_dimensions.items()
        if dim.customWidth and dim.width
    }
    grid.freeze_panes = ws.freeze_panes
    return grid


def _load_calamine(path: PathLike) -> WorkbookModel:
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(os.fspath(path))
    sheets = []
    for name in wb.sheet_names:
        rows = wb.get_sheet_by_name(name).to_python(skip_empty_area=False)
        # calamine hands back "" for empty cells, and floats for every number
        cleaned = [[_from_calamine(v) for v in row] for row in rows]
        while cleaned and all(v is None for v in cleaned[-1]):
            cleaned.pop()
        sheets.append(SheetGrid.from_values(name, cleaned))
    return WorkbookModel(sheets)


def _from_calamine(v: Any) -> Any:
    if isinstance(v, str) and v == "":
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def dump_xlsx(model: WorkbookModel, path: PathLike) -> None:
    """Write a :class:`~.model.WorkbookModel` to an ``.xlsx`` file with openpyxl.

    Formula cells are written as formulas, so their cached results are lost: openpyxl
    has no calculation engine, and Excel recomputes them on open.

    Parameters
    ----------
    model : WorkbookModel
        What to write.
    path : str or os.PathLike
        Where to write it.
    """
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for grid in model.sheets:
        ws = wb.create_sheet(grid.name)
        _write_sheet(ws, grid, model.number_formats)
    for name, text in model.defined_names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=text)
    wb.save(path)


def _write_sheet(ws: Any, grid: SheetGrid, numfmts: NumberFormats) -> None:
    kind = grid.arrays[Slot.KIND]
    formula = grid.arrays[Slot.FORMULA]
    numfmt = grid.arrays[Slot.NUMFMT]
    values = grid.values()
    occupied = (kind != CellKind.EMPTY) | (formula != "") | (numfmt != 0)
    for i, j in zip(*np.nonzero(occupied), strict=True):
        f = str(formula[i, j])
        value = f if f else values[i, j]
        cell = ws.cell(row=int(i) + 1, column=int(j) + 1)
        cell.value = value
        if kind[i, j] == CellKind.STRING and not f:
            # a string that happens to start with "=" is still a string: openpyxl's
            # value setter would have bound it as a formula
            cell.data_type = "s"
        if numfmt[i, j]:
            cell.number_format = numfmts[int(numfmt[i, j])]
    for rng in grid.merged:
        ws.merge_cells(rng)
    for letter, width in grid.col_widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = grid.freeze_panes
