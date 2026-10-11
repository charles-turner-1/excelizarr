"""Shared fixtures and the hypothesis profile.

Most properties write a real ``.xlsx`` through openpyxl on every example, which is
tens of milliseconds a time, so the per-example deadline is off: a slow example is not
a failure here, and openpyxl's timing jitter would otherwise make the suite flaky.
"""

import datetime as dt
import os
from pathlib import Path

import openpyxl
import pytest
from hypothesis import HealthCheck, settings

settings.register_profile(
    "default",
    deadline=None,
    max_examples=60,
    suppress_health_check=[HealthCheck.too_slow, HealthCheck.function_scoped_fixture],
)
settings.register_profile(
    "ci", parent=settings.get_profile("default"), max_examples=200
)
settings.load_profile(os.environ.get("HYPOTHESIS_PROFILE", "default"))


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    """A small hand-made workbook with a bit of everything.

    Returns
    -------
    pathlib.Path
        ``Inputs`` (a header row, mixed types, a formula, a merge, frozen panes, a
        percentage format), ``Empty`` (nothing at all) and ``Lookup`` (a two-column
        table), plus a workbook-scoped defined name.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(["name", "x", "when", "ok"])
    ws.append(["a", 1, dt.datetime(2024, 1, 2, 3, 4, 5), True])
    ws.append(["b", 2.5, dt.datetime(2025, 6, 7), False])
    ws["E2"] = "=B2*2"
    ws["B3"].number_format = "0.00%"
    ws.merge_cells("G1:H2")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 30
    wb.create_sheet("Empty")
    lookup = wb.create_sheet("Lookup")
    for row in [("key", "value"), ("alpha", 1.5), ("beta", 2.5)]:
        lookup.append(row)
    from openpyxl.workbook.defined_name import DefinedName

    wb.defined_names["rates"] = DefinedName("rates", attr_text="Lookup!$B$2:$B$3")
    path = tmp_path / "sample.xlsx"
    wb.save(path)
    return path
