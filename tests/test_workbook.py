"""The lazy workbook API, by example: the parts the properties don't reach."""

import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
import pytest
import zarr

import excelizarr as xz
from excelizarr.io_zarr import sheet_key


@pytest.fixture
def wb(sample_xlsx: Path) -> xz.ZarrWorkbook:
    """The sample workbook, converted to an in-memory store."""
    return xz.xlsx_to_zarr(sample_xlsx, zarr.storage.MemoryStore())


def test_sheets_in_tab_order(wb: xz.ZarrWorkbook) -> None:
    """Sheet names, positions and iteration all follow tab order."""
    assert wb.sheet_names == ["Inputs", "Empty", "Lookup"]
    assert wb[1].name == "Empty"
    assert [s.name for s in wb] == wb.sheet_names
    assert len(wb) == 3
    assert "Inputs" in repr(wb)


def test_cells_ranges_and_types(wb: xz.ZarrWorkbook) -> None:
    """Values come back as the Python types openpyxl would give."""
    s = wb["Inputs"]
    assert s.shape == (3, 8)
    assert s.dimensions == "A1:H3"
    assert s["A2"] == "a"
    assert s["B2"] == 1 and isinstance(s["B2"], int)
    assert s["C2"] == dt.datetime(2024, 1, 2, 3, 4, 5)
    assert s["D3"] is False
    assert s["Z99"] is None
    np.testing.assert_array_equal(
        wb["Inputs!A1:B2"], np.array([["name", "x"], ["a", 1]], dtype=object)
    )


def test_formulas_formats_and_layout(wb: xz.ZarrWorkbook) -> None:
    """Formula text, number formats, merges and defined names all survive."""
    s = wb["Inputs"]
    assert s.formulas("E2")[0, 0] == "=B2*2"
    assert s.formulas("A1")[0, 0] is None
    assert s.number_formats("B3")[0, 0] == "0.00%"
    assert s.merged == ["G1:H2"]
    assert wb.defined_names == {"rates": "Lookup!$B$2:$B$3"}


def test_empty_sheet(wb: xz.ZarrWorkbook) -> None:
    """A sheet with nothing in it is ``0x0`` and reads as empty."""
    s = wb["Empty"]
    assert s.shape == (0, 0)
    assert s["A1"] is None
    assert s["A1:C3"].shape == (0, 0)


def test_back_to_xlsx(wb: xz.ZarrWorkbook, tmp_path: Path) -> None:
    """``zarr_to_xlsx`` writes a file Excel (well, openpyxl) can open, layout intact."""
    out = tmp_path / "back.xlsx"
    xz.zarr_to_xlsx(wb, out)
    ws = openpyxl.load_workbook(out)["Inputs"]
    assert ws["E2"].value == "=B2*2"
    assert ws["B3"].number_format == "0.00%"
    assert ws.freeze_panes == "B2"
    assert ws.column_dimensions["A"].width == 30
    assert [str(r) for r in ws.merged_cells.ranges] == ["G1:H2"]


def test_in_place_writes(sample_xlsx: Path, tmp_path: Path) -> None:
    """``r+`` stores take cell and block writes, and formulas by their leading ``=``."""
    store = tmp_path / "w.zarr"
    xz.xlsx_to_zarr(sample_xlsx, store)
    s = xz.open_workbook(store, mode="r+")["Inputs"]
    s["B2"] = 99
    s["A5:B6"] = [[1, 2], [3, "four"]]
    s["F1"] = "=SUM(B2:B3)"
    again = xz.open_workbook(store)["Inputs"]
    assert again["B2"] == 99
    assert again["A5:B6"].tolist() == [[1, 2], [3, "four"]]
    assert again.formulas("F1")[0, 0] == "=SUM(B2:B3)"
    assert again["F1"] is None  # no cached value, as with any formula openpyxl wrote
    with pytest.raises(xz.InvalidReferenceError):
        s["A:B"] = 1


def test_read_only_by_default(
    wb: xz.ZarrWorkbook, sample_xlsx: Path, tmp_path: Path
) -> None:
    """``open_workbook`` opens read-only unless asked otherwise."""
    store = tmp_path / "ro.zarr"
    xz.xlsx_to_zarr(sample_xlsx, store)
    with pytest.raises(Exception):  # zarr's own read-only error
        xz.open_workbook(store)["Inputs"]["A1"] = "nope"


def test_refuses_to_overwrite(sample_xlsx: Path, tmp_path: Path) -> None:
    """The default mode won't clobber an existing store; ``mode="w"`` will."""
    store = tmp_path / "x.zarr"
    xz.xlsx_to_zarr(sample_xlsx, store)
    with pytest.raises(Exception):
        xz.xlsx_to_zarr(sample_xlsx, store)
    xz.xlsx_to_zarr(sample_xlsx, store, mode="w")


def test_not_a_workbook() -> None:
    """A plain zarr group is not a workbook."""
    store = zarr.storage.MemoryStore()
    zarr.open_group(store, mode="w")
    with pytest.raises(xz.FormatVersionError):
        xz.open_workbook(store)


def test_newer_format_is_refused(wb: xz.ZarrWorkbook) -> None:
    """A store from a future format version is refused rather than misread."""
    meta = dict(wb.root.attrs["excelizarr"])
    meta["format"] = 99
    root = zarr.open_group(wb.root.store, mode="r+")
    root.attrs["excelizarr"] = meta
    with pytest.raises(xz.FormatVersionError, match="format 99"):
        xz.open_workbook(wb.root.store)


def test_missing_sheet_is_a_key_error(wb: xz.ZarrWorkbook) -> None:
    """Asking for a sheet or range that doesn't exist is a ``KeyError``."""
    for key in ["Nope", 5, "Nope!A1", "A1", "not a ref at all"]:
        with pytest.raises(KeyError):
            wb[key]


def test_zip_store(sample_xlsx: Path, tmp_path: Path) -> None:
    """A zip of XML becomes a zip of chunks: ``ZipStore`` works like any other store."""
    path = tmp_path / "book.zarr.zip"
    with zarr.storage.ZipStore(path, mode="w") as store:
        xz.xlsx_to_zarr(sample_xlsx, store)
    with zarr.storage.ZipStore(path, mode="r") as store:
        assert xz.open_workbook(store)["Lookup"]["A2"] == "alpha"


@pytest.mark.parametrize(
    ("name", "key"),
    [
        ("Sheet1", "Sheet1"),
        ("a/b", "a%2Fb"),
        ("100%", "100%25"),
        ("__x", "%5F_x"),
        (".hid", "%2Ehid"),
    ],
)
def test_sheet_keys(name: str, key: str) -> None:
    """Sheet names escape to legal, still-readable zarr node names."""
    assert sheet_key(name) == key
